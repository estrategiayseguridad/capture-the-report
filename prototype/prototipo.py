import argparse
from docx import Document
from docx.shared import Pt
import pandas as pd


def generar_reporte_word(csv_filepath, docx_output_path):
    # 1. Cargar el archivo CSV
    df = pd.read_csv(csv_filepath)

    # 2. Calcular los totales y desgloses
    total_alertas = len(df)

    severidades = df['kibana.alert.severity'].str.lower().value_counts()
    altas = severidades.get('high', 0)
    medias = severidades.get('medium', 0)
    bajas = severidades.get('low', 0)

    reglas_top = df['kibana.alert.rule.name'].value_counts()

    # 3. Crear el documento Word
    doc = Document()

    # Formato general
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)

    # Parrafo total
    p_total = doc.add_paragraph()
    p_total.add_run(f'{total_alertas} alertas').bold = True
    p_total.add_run(' se monitorearon en total.')

    # Listado por severidad
    p_altas = doc.add_paragraph(style='List Bullet')
    p_altas.add_run(f'{altas} ').bold = True
    p_altas.add_run('alertas fueron ')
    p_altas.add_run('altas.').bold = True

    p_medias = doc.add_paragraph(style='List Bullet')
    p_medias.add_run(f'{medias} ').bold = True
    p_medias.add_run('alerta fue ')
    p_medias.add_run('media.').bold = True

    p_bajas = doc.add_paragraph(style='List Bullet')
    p_bajas.add_run(f'{bajas} ').bold = True
    p_bajas.add_run('alertas fueron ')
    p_bajas.add_run('bajas.').bold = True

    doc.add_paragraph()

    # Titulo reglas
    p_reglas_titulo = doc.add_paragraph()
    p_reglas_titulo.add_run(
        'Las reglas que más alertas generaron fueron:'
    ).bold = True

    for regla, cantidad in reglas_top.items():
        p_regla = doc.add_paragraph()
        p_regla.add_run(f'{regla}\t')
        p_regla.add_run(str(cantidad)).bold = True

    # 4. Guardar archivo
    doc.save(docx_output_path)
    print(f'✅ Reporte generado exitosamente en: {docx_output_path}')


if __name__ == '__main__':
    # Configuración de flags de la terminal
    parser = argparse.ArgumentParser(
        description='Genera un reporte de Word a partir de un archivo CSV de alertas.'
    )

    # Flag de entrada (-i o --input)
    parser.add_argument(
        '-i',
        '--input',
        type=str,
        required=True,
        help='Ruta del archivo CSV de entrada',
    )

    # Flag opcional de salida (-o o --output)
    parser.add_argument(
        '-o',
        '--output',
        type=str,
        default='reporte_alertas.docx',
        help='Ruta/Nombre del archivo de Word a generar (opcional)',
    )

    args = parser.parse_args()

    # Ejecutar función con las flags recibidas
    generar_reporte_word(args.input, args.output)